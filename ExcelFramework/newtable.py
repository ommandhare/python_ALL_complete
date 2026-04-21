import pandas as pd
import formatConfig as fcg
from openpyxl import Workbook
from openpyxl.chart import PieChart,BarChart,BarChart3D, Reference
import numpy as np
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


templatePath = r"C:\Users\Om Mandhare\PycharmProjects\python_ALL_complete\ExcelFramework\template.csv"
path = r"C:\Users\Om Mandhare\PycharmProjects\python_ALL_complete\ExcelFramework\product.csv"
outPath = r"C:\Users\Om Mandhare\PycharmProjects\python_ALL_complete\ExcelFramework\tableproduct.xlsx"

df = pd.read_csv(path)
#print(df.values)
varValDict = {}
count=0
for line in open(templatePath):
    if(count==0):
        count+=1
        continue
    print(line)
    file,wsheet,obj,oName,variable,value, type = line.strip().split(",")
    key = file +"_" + wsheet + "_" + oName + "_" + variable
    vValue = value
    if(type=='i'):
        varValDict[key] = int(vValue)
    else:
        varValDict[key] = vValue
print(varValDict)

df_agg = df.groupby("category")["price"].sum().reset_index()

with pd.ExcelWriter(outPath, engine='xlsxwriter') as writer:
    Data_sheet = "Sheet1"
    Chart_sheet = "chart"
    
    df.to_excel(writer, sheet_name=Data_sheet, index=False, startrow=2)

    df_agg.to_excel(writer, sheet_name=Chart_sheet, index=False, startcol=1)


    workbook = writer.book
    data_ws = writer.sheets[Data_sheet]
    chart_ws =  writer.sheets[Chart_sheet]

    start_row = 2
    start_col = 0
    end_row = start_row + len(df)
    end_col = start_col + len(df.columns) - 1

    data_ws.add_table(
        start_row,
        start_col,
        end_row,
        end_col,
        {
            "style": fcg.tableStyleDict[varValDict['f1_w1_sec2_tableStyle']],
            "columns": [{"header": col} for col in df.columns],# <--- STYLE FROM formatConfig
            "autofilter":False
        },
    )

    print("Excel with table format created at:", outPath)

    # =====================================================
    # SIMPLE BAR CHART (XlsxWriter)
    # =====================================================

    chart = workbook.add_chart({'type': 'column'})

    # Data boundaries
    first_data_row = start_row + 1  # first row of actual data (after header)
    last_data_row = end_row  # last row of data

    name_col = start_col + 1  # Name column index
    value_col = start_col + 2  # Price column index (0=ID,1=Name,2=price)

    # Add series
    chart.add_series({
        'name': 'Price',
        'categories': [
            data_ws.name,
            first_data_row, name_col,
            last_data_row, name_col
        ],
        'values': [
            data_ws.name,
            first_data_row, value_col,
            last_data_row, value_col
        ]
    })

    # Chart titles (optional)
    chart.set_title({'name': 'Price Chart'})
    chart.set_x_axis({'name': 'Product Name'})
    chart.set_y_axis({'name': 'Price'})

    # Insert chart in sheet
    data_ws.insert_chart("G18", chart)

   # Aggregate




    pie = workbook.add_chart({'type': 'pie'})

    start_row = 1
    end_row   = start_row + (len(df_agg)-1)

    categories_col = 1   # aggregated category
    values_col= 2   # aggregated value (price or qty)

    pie.add_series({
        'Category': 'Category Distribution',
        'categories': [Chart_sheet, start_row , categories_col, end_row, categories_col],
        'values':     [Chart_sheet, start_row , values_col, end_row, values_col],
        'points':   fcg.pieColorSets[varValDict['f1_w1_pie1_pie']]['points']
    })

    pie.set_title({'name': "Aggregated Pie Chart"})
    pie.set_style(10)

    chart_ws.insert_chart("G3", pie)


    #Pareto


    pareto_col = "qty"

    df_pareto = df.groupby("category")[pareto_col].sum().reset_index()
    df_pareto = df_pareto.sort_values(by=pareto_col, ascending=True)

    df_pareto["cum_sum"] = df_pareto[pareto_col].cumsum()
    df_pareto["cum_pct"] = 100 * df_pareto["cum_sum"] / df_pareto[pareto_col].sum()

    pareto_start = 25
    first_data_row = 1
    df_pareto.to_excel(writer, sheet_name=Data_sheet, index=False, startcol=pareto_start)

    bar_chart = workbook.add_chart({'type': 'column'})
    line_chart = workbook.add_chart({'type': 'line'})

    start = pareto_start
    end = start + len(df_pareto)
    print(len(df_pareto))
    print(start)
    print(end)

    cat_col = pareto_start
    value_col = pareto_start+1
    cum_pct_col = 28 # cum_pct


    # Bar Chart (Quantity)
    bar_chart.add_series({
        'name': 'Qty',
        'categories': [Data_sheet, first_data_row, cat_col, len(df_pareto), cat_col],
        'values': [Data_sheet, first_data_row, value_col, len(df_pareto), value_col],
    })



    line_chart.add_series({
        'name': 'Cumulative %',
        'categories': [Data_sheet, first_data_row, cat_col, len(df_pareto), cat_col],
        'values': [Data_sheet, first_data_row, cum_pct_col, len(df_pareto), cum_pct_col],
        'y2_axis': True,
        'marker': {'type': 'circle'},
    })

    # Combine charts
    bar_chart.combine(line_chart)

    bar_chart.set_title({'name': 'Pareto Chart'})
    bar_chart.set_x_axis({'name': 'Category'})
    bar_chart.set_y_axis({'name': 'Qty'})
    bar_chart.set_y2_axis({'name': 'Cumulative %'})

    data_ws.insert_chart("Y16", bar_chart)




